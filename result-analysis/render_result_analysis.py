#!/usr/bin/env python3
"""Render the RESULT ANALYSIS FORMAT proforma for each class, from the verified
Half Yearly CSVs, as both a printable PDF (reportlab) and an editable XLSX
(openpyxl) - 9-column grid extended to 10 with the added 75%-90% band.
"""

import csv
import os

import pymupdf
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))

SCHOOL = "DAV PUBLIC SCHOOL, CANTT. AREA, GAYA JEE"
COMMITTEE = "(Managed By \u2013 DAV College Managing Committee, New Delhi \u2013 55)"
TITLE = "Result Analysis"

EXAM_OPTIONS = [
    ("PT-1/2/3/4", ("PT-1", "PT-2", "PT-3", "PT-4")),
    ("HALF \u2013 YEARLY", ("HALF-YEARLY", "HALF YEARLY")),
    ("PRE - BOARD", ("PRE-BOARD", "PRE BOARD")),
    ("ANNUAL", ("ANNUAL",)),
]

BAND_LABELS = [
    "Less than 40%", "40% - 50%", "50% - 60%",
    "60% \u2013 75 %", "75% - 90%", "More than 90%",
]
BAND_KEYS = [
    "less_than_40", "40_to_50", "50_to_60", "60_to_75", "75_to_90", "more_than_90",
]

SUBJECTS = [
    "English", "Hindi", "Sanskrit", "Mathematics", "Science", "Physics",
    "Chemistry", "Biology", "Social Sc.", "Accountancy", "B. St.",
    "Economics", "Physical Edu.", "A.I.",
]
SPARE_ROWS = 1
N_ROWS = len(SUBJECTS) + SPARE_ROWS

CLASSES = ["IX A", "VII A", "X C", "X F", "X H", "X I", "XI D"]

# column widths, total = 523 pt (A4 595.32 - 2 x 36 margin)
COLW = [76, 46, 50, 38, 38, 38, 38, 38, 38, 123]
MARGIN = 36


def load_rows(section, exam="HALF-YEARLY"):
    path = os.path.join(ROOT, section, "Result Analysis - %s 2026-27.csv" % exam)
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    return rows


def ticked(exam_type):
    return [
        "[X]" if exam_type in keys else "[ ]"
        for _, keys in EXAM_OPTIONS
    ]


def exam_line(exam_type):
    marks = ticked(exam_type)
    parts = ["%s %s" % (m, label) for m, (label, _) in zip(marks, EXAM_OPTIONS)]
    return "NAME OF EXAM (PLEASE WRITE OR TICK):  " + "      ".join(parts)


def build_pdf(section, rows, out_path):
    first = rows[0]
    exam_type = first["exam_type"]
    styles = {
        "school": ParagraphStyle("school", fontName="Helvetica-Bold", fontSize=14,
                                 alignment=TA_CENTER, leading=17),
        "committee": ParagraphStyle("committee", fontName="Helvetica", fontSize=9,
                                    alignment=TA_CENTER, leading=12),
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=14,
                                alignment=TA_CENTER, leading=18, spaceBefore=8),
        "exam": ParagraphStyle("exam", fontName="Helvetica-Bold", fontSize=9,
                               leading=13),
        "detail": ParagraphStyle("detail", fontName="Helvetica", fontSize=9.5,
                                 leading=15),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=9.5,
                               alignment=TA_CENTER, leading=11.5),
        "subj": ParagraphStyle("subj", fontName="Helvetica", fontSize=9.5,
                               leading=11.5),
        "hcell": ParagraphStyle("hcell", fontName="Helvetica-Bold", fontSize=9.5,
                                alignment=TA_CENTER, leading=11.5),
        "band": ParagraphStyle("band", fontName="Helvetica-Bold", fontSize=7.2,
                               alignment=TA_CENTER, leading=8.5),
        "box": ParagraphStyle("box", fontName="Helvetica-Bold", fontSize=9,
                              alignment=TA_CENTER, leading=13),
        "sig": ParagraphStyle("sig", fontName="Helvetica-Bold", fontSize=9.5,
                              alignment=TA_CENTER, leading=13),
        "small": ParagraphStyle("small", fontName="Helvetica", fontSize=9.5,
                                leading=15),
    }

    story = [
        Paragraph(SCHOOL, styles["school"]),
        Paragraph(COMMITTEE, styles["committee"]),
        Paragraph(TITLE, styles["title"]),
        Spacer(1, 12),
        Paragraph(exam_line(exam_type), styles["exam"]),
        Spacer(1, 8),
        Paragraph(
            "Class &amp; Section: <b>%s</b>"
            "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Period of Examination: _________ to _________"
            % first["class_section"],
            styles["detail"]),
        Paragraph(
            "Class Strength: <b>%s</b>"
            "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Class Teacher's Name: <b>%s</b>"
            % (first["class_strength"],
               first["class_teacher_name"] or "________________"),
            styles["detail"]),
        Spacer(1, 10),
    ]

    # ---- 10-column grid -------------------------------------------------
    header = [
        [Paragraph("Subject", styles["hcell"]),
         Paragraph("Highest Marks.", styles["hcell"]),
         Paragraph("Average Marks", styles["hcell"]),
         Paragraph("No. of students who got", styles["hcell"]), "", "", "", "", "",
         Paragraph("Name of Subject Teacher", styles["hcell"])],
        ["", "", ""] + [Paragraph(x, styles["band"]) for x in BAND_LABELS] + [""],
    ]

    body = []
    for row in rows[:N_ROWS]:
        subject = row["subject"] or ""
        cells = [Paragraph(subject, styles["subj"])]
        cells.append(Paragraph(row["highest_marks"] or "", styles["cell"]))
        cells.append(Paragraph(row["average_marks"] or "", styles["cell"]))
        for key in BAND_KEYS:
            cells.append(Paragraph(row[key] or "", styles["cell"]))
        cells.append(Paragraph(row["subject_teacher"] or "", styles["cell"]))
        body.append(cells)

    data = header + body
    # fill the page: body rows grow so the table + totals + signature
    # bottom out at the bottom margin instead of leaving a blank void.
    # The row height is self-calibrated against the rendered page so the
    # report always fits one A4 page regardless of installed fonts.
    style = [
        # outer + inner grid
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("BOX", (0, 0), (-1, -1), 1.1, colors.black),
        # merged header cells
        ("SPAN", (3, 0), (8, 0)),
        ("SPAN", (0, 0), (0, 1)),
        ("SPAN", (1, 0), (1, 1)),
        ("SPAN", (2, 0), (2, 1)),
        ("SPAN", (9, 0), (9, 1)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 2), (-1, -1), "CENTER"),
        ("BACKGROUND", (0, 0), (-1, 1), colors.Color(0.93, 0.93, 0.93)),
    ]

    totals = Table(
        [[Paragraph("No. of Students Passed:", styles["detail"]),
          Paragraph(first["no_students_passed"], styles["box"]),
          Paragraph("Detained:", styles["detail"]),
          Paragraph(first["detained"], styles["box"]),
          Paragraph("Advised for Improvement:", styles["detail"]),
          Paragraph(first["advised_for_improvement"], styles["box"])]],
        colWidths=[140, 44, 120, 44, 131, 44], rowHeights=[26],
    )
    totals.setStyle(TableStyle([
        ("BOX", (1, 0), (1, 0), 1.1, colors.black),
        ("BOX", (3, 0), (3, 0), 1.1, colors.black),
        ("BOX", (5, 0), (5, 0), 1.1, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
    ]))

    sig = Table(
        [["", "", ""],
         [Paragraph("Date: ______________", styles["small"]),
          Paragraph("Class Teacher", styles["sig"]),
          Paragraph("PRINCIPAL", styles["sig"])]],
        colWidths=[175, 174, 174], rowHeights=[24, 18],
    )
    sig.setStyle(TableStyle([
        ("LINEABOVE", (1, 0), (1, 0), 0.8, colors.black),
        ("LINEABOVE", (2, 0), (2, 0), 0.8, colors.black),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
    ]))

    frame_h = A4[1] - 2 * MARGIN
    hh0, hh1 = 30, 22

    def _build(body_h):
        body_hts = [int(body_h)] * N_ROWS
        for i in range(round((body_h - int(body_h)) * N_ROWS)):
            body_hts[i] += 1
        table = Table(data, colWidths=COLW, rowHeights=[hh0, hh1] + body_hts)
        table.setStyle(TableStyle(style))
        full = list(story) + [table, Spacer(1, 12), totals, Spacer(1, 26), sig]
        doc = SimpleDocTemplate(
            out_path, pagesize=A4,
            leftMargin=MARGIN, rightMargin=MARGIN,
            topMargin=MARGIN, bottomMargin=MARGIN,
            title="%s %s - %s" % (TITLE, first["class_section"], first["exam_name"]),
            author="DAV Public School, Cantt. Area, Gaya",
        )
        doc.build(full)

    def _measure(path):
        """(pages, blank_from_edge) of the rendered report."""
        d = pymupdf.open(path)
        p = d[d.page_count - 1]
        drawn = [it["rect"][3] for it in p.get_drawings()]
        tbot = max((w[3] for w in p.get_text("words")), default=0)
        bottom = max([tbot] + drawn)
        blank = (MARGIN + frame_h) - bottom - (d.page_count - 1) * frame_h
        return d.page_count, blank

    # body height is monotonic: 1 page when small, then overflow.  Binary
    # search for the largest body that keeps everything on one page, then
    # polish it to leave a few points of bottom margin.
    lo, hi = 8.0, None
    best = None
    for _ in range(16):
        b = lo * 2 if hi is None else (lo + hi) / 2
        _build(b)
        pages, blank = _measure(out_path)
        if pages == 1:
            lo, best = b, (b, blank)
            if 5.0 <= blank <= 14.0:
                return
        else:
            hi = b
    if best:
        b, _blank = best
        _build(b)


THIN = Side(style="thin", color="000000")
MED = Side(style="medium", color="000000")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX = Border(left=MED, right=MED, top=MED, bottom=MED)


def build_xlsx(section, rows, out_path):
    first = rows[0]
    exam_type = first["exam_type"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result Analysis"
    ws.sheet_view.showGridLines = False

    widths = [22, 13, 14, 11, 11, 11, 11, 11, 11, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last_col = get_column_letter(len(widths))

    def merge_text(row, text, size, bold=True, italic=False, height=None):
        ws.merge_cells("A%d:%s%d" % (row, last_col, row))
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(name="Calibri", size=size, bold=bold, italic=italic)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if height:
            ws.row_dimensions[row].height = height

    merge_text(1, SCHOOL, 15, height=22)
    merge_text(2, COMMITTEE, 10, bold=False, height=16)
    merge_text(3, TITLE, 15, height=24)
    ws.row_dimensions[4].height = 8

    # exam selector
    ws.merge_cells("A5:%s5" % last_col)
    c = ws.cell(row=5, column=1,
                value="NAME OF EXAM (PLEASE WRITE OR TICK):")
    c.font = Font(bold=True, size=10)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[5].height = 18

    marks = ticked(exam_type)
    col = 1
    for mark, (label, _) in zip(marks, EXAM_OPTIONS):
        cell = ws.cell(row=6, column=col, value="%s %s" % (mark, label))
        cell.font = Font(size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        col += 3
    ws.row_dimensions[6].height = 18

    details = [
        (7, "Class & Section: %s" % first["class_section"],
         "Period of Examination: _______ to _______"),
        (8, "Class Strength: %s" % first["class_strength"],
         "Class Teacher's Name: %s" % (first["class_teacher_name"] or "__________")),
    ]
    for row, left, right in details:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cl = ws.cell(row=row, column=1, value=left)
        cl.font = Font(bold=True, size=10)
        cl.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=row, start_column=5, end_row=row,
                       end_column=len(widths))
        cr = ws.cell(row=row, column=5, value=right)
        cr.font = Font(size=10)
        cr.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18

    ws.row_dimensions[9].height = 8

    # ---- grid -----------------------------------------------------------
    h1, h2 = 10, 11
    spans = [(1, 1, 1, 2), (2, 1, 2, 2), (3, 1, 3, 2), (4, 1, 9, 1), (10, 1, 10, 2)]
    for c1, r1, c2, r2 in spans:
        ws.merge_cells(start_row=h1 + r1 - 1, start_column=c1,
                       end_row=h1 + r2 - 1, end_column=c2)

    ws.cell(row=h1, column=1, value="Subject")
    ws.cell(row=h1, column=2, value="Highest Marks.")
    ws.cell(row=h1, column=3, value="Average Marks")
    ws.cell(row=h1, column=4, value="No. of students who got")
    ws.cell(row=h1, column=10, value="Name of Subject Teacher")
    for i, label in enumerate(BAND_LABELS):
        ws.cell(row=h2, column=4 + i, value=label)

    for row in range(h1, h2 + 1):
        ws.row_dimensions[row].height = 30 if row == h1 else 26
        for col in range(1, len(widths) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            cell.border = GRID
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="EDEDED")

    data_row = h2 + 1
    for index, row in enumerate(rows[:N_ROWS]):
        r = data_row + index
        ws.row_dimensions[r].height = 18
        values = [row["subject"], row["highest_marks"], row["average_marks"]]
        values += [row[k] for k in BAND_KEYS]
        values += [row["subject_teacher"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col,
                           value=("" if value == "" else
                                  int(value) if str(value).isdigit() else value))
            cell.border = GRID
            cell.alignment = Alignment(
                horizontal="left" if col in (1, 10) else "center",
                vertical="center")
            cell.font = Font(size=10)

    last_grid = data_row + N_ROWS - 1

    # ---- totals ---------------------------------------------------------
    t = last_grid + 2
    ws.merge_cells(start_row=t, start_column=1, end_row=t, end_column=5)
    c = ws.cell(row=t, column=1, value="No. of Students Passed:")
    c.font = Font(bold=True, size=11)
    c.alignment = Alignment(horizontal="right", vertical="center")
    b = ws.cell(row=t, column=6, value=int(first["no_students_passed"]))
    b.font = Font(bold=True, size=11)
    b.alignment = Alignment(horizontal="center", vertical="center")
    b.border = BOX

    ws.merge_cells(start_row=t, start_column=7, end_row=t, end_column=9)
    c = ws.cell(row=t, column=7, value="Detained:")
    c.font = Font(bold=True, size=11)
    c.alignment = Alignment(horizontal="right", vertical="center")
    b = ws.cell(row=t, column=10, value=int(first["detained"]))
    b.font = Font(bold=True, size=11)
    b.alignment = Alignment(horizontal="center", vertical="center")
    b.border = BOX
    ws.row_dimensions[t].height = 24

    t2 = t + 1
    ws.merge_cells(start_row=t2, start_column=1, end_row=t2, end_column=5)
    c = ws.cell(row=t2, column=1, value="Advised for Improvement:")
    c.font = Font(bold=True, size=11)
    c.alignment = Alignment(horizontal="right", vertical="center")
    b = ws.cell(row=t2, column=6, value=int(first["advised_for_improvement"]))
    b.font = Font(bold=True, size=11)
    b.alignment = Alignment(horizontal="center", vertical="center")
    b.border = BOX
    ws.row_dimensions[t2].height = 24

    # ---- sign off -------------------------------------------------------
    s = t2 + 3
    ws.row_dimensions[s].height = 22
    d = ws.cell(row=s, column=1, value="Date: ______________")
    d.font = Font(size=11)
    ws.merge_cells(start_row=s, start_column=6, end_row=s, end_column=7)
    ct = ws.cell(row=s, column=6, value="Class Teacher")
    ct.font = Font(bold=True, size=11)
    ct.alignment = Alignment(horizontal="center")
    ct.border = Border(top=MED)
    ws.merge_cells(start_row=s, start_column=9, end_row=s, end_column=10)
    pr = ws.cell(row=s, column=9, value="PRINCIPAL")
    pr.font = Font(bold=True, size=11)
    pr.alignment = Alignment(horizontal="center")
    pr.border = Border(top=MED)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:%s%d" % (last_col, s)

    wb.save(out_path)


def main():
    made = []
    for section in CLASSES:
        folder = os.path.join(ROOT, section)
        rows = load_rows(section)
        stem = os.path.join(folder, "Result Analysis - HALF-YEARLY 2026-27")
        build_pdf(section, rows, stem + ".pdf")
        build_xlsx(section, rows, stem + ".xlsx")
        made += [stem + ".pdf", stem + ".xlsx"]
        print("%-6s ok  strength=%s passed=%s detained=%s improve=%s"
              % (section, rows[0]["class_strength"], rows[0]["no_students_passed"],
                 rows[0]["detained"], rows[0]["advised_for_improvement"]))

    print("\n%d file(s) written." % len(made))
    for path in made:
        print("  " + os.path.relpath(path, ROOT))


if __name__ == "__main__":
    main()
